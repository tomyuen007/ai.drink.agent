"""
Render App.tsx UI in iOS, Android, and Web device frames and embed into ui.docx.
"""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import datetime

ROOT = Path(__file__).parent
DOCX = ROOT / "ui.docx"
F    = "/usr/share/fonts/truetype/ubuntu"

# ── Font helper ────────────────────────────────────────────────────────────────
def font(size, bold=False):
    try:
        return ImageFont.truetype(f"{F}/Ubuntu[wdth,wght].ttf", size)
    except Exception:
        return ImageFont.load_default()

# ── Color palette ──────────────────────────────────────────────────────────────
C = {
    "bg":        "#E0F2FE", "white":    "#FFFFFF",
    "sky50":     "#F0F9FF", "sky200":   "#BAE6FD",
    "sky300":    "#7DD3FC", "sky500":   "#0EA5E9",
    "sky600":    "#0284C7", "sky700":   "#0369A1",
    "slate500":  "#64748B",
    "gray300":   "#D1D5DB", "gray400":  "#9CA3AF",
    "gray600":   "#4B5563", "gray700":  "#374151",
    "gray900":   "#111827",
    "green50":   "#F0FDF4", "green300": "#86EFAC", "green600": "#16A34A",
    "red50":     "#FEF2F2", "red300":   "#FCA5A5", "red700":   "#B91C1C",
    "violet100": "#EDE9FE", "violet200":"#DDD6FE", "violet700": "#6D28D9",
}

def rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

def rr(draw, xy, radius, fill, outline=None, ow=2):
    draw.rounded_rectangle(xy, radius=radius, fill=rgb(fill),
                            outline=rgb(outline) if outline else None, width=ow)

def tw(draw, txt, fnt):
    bb = draw.textbbox((0, 0), txt, font=fnt)
    return bb[2] - bb[0]


# ══════════════════════════════════════════════════════════════════════════════
# APP CONTENT RENDERER
# ══════════════════════════════════════════════════════════════════════════════
def _render_content(cw: int, show_answer: bool) -> Image.Image:
    """Draw app UI at content width cw. Returns a cropped image exactly as tall as the content."""
    pad   = int(cw * 0.065)
    inner = cw - pad * 2
    y     = int(cw * 0.13)
    scale = cw / 390

    def s(v):  return max(1, int(v * scale))
    def fs(v): return max(8, int(v * scale))

    img  = Image.new("RGB", (cw, 2400), rgb(C["bg"]))
    draw = ImageDraw.Draw(img)

    # Header
    tf = font(fs(22), bold=True)
    sf = font(fs(11))
    t  = "Weather AI"
    draw.text(((cw - tw(draw, t, tf)) // 2, y), t, font=tf, fill=rgb(C["sky700"]))
    y += s(34)
    sub = "Claude · MCP · RAG · Open-Meteo"
    draw.text(((cw - tw(draw, sub, sf)) // 2, y), sub, font=sf, fill=rgb(C["slate500"]))
    y += s(36)

    lf  = font(fs(12))
    inf = font(fs(13))

    def labeled_input(label, value, placeholder, yp):
        draw.text((pad, yp), label, font=lf, fill=rgb(C["gray700"]))
        yp += s(20)
        h = s(42)
        rr(draw, [pad, yp, pad + inner, yp + h], s(10), C["white"], C["gray300"])
        txt = value or placeholder
        col = C["gray900"] if value else C["gray400"]
        draw.text((pad + s(12), yp + s(12)), txt, font=inf, fill=rgb(col))
        return yp + h + s(14)

    y = labeled_input("City", "", "e.g. Toronto, Tokyo, London", y)
    y = labeled_input("Question", "Should I bring an umbrella?", "", y)

    # Quick questions dropdown
    draw.text((pad, y), "Quick questions", font=lf, fill=rgb(C["gray700"]))
    y += s(20)
    dh = s(44)
    rr(draw, [pad, y, pad + inner, y + dh], s(10), C["white"], C["gray300"])
    draw.text((pad + s(12), y + s(13)), "Should I bring an umbrella?", font=inf, fill=rgb(C["gray900"]))
    draw.text((pad + inner - s(24), y + s(12)), "▾", font=font(fs(14)), fill=rgb(C["gray400"]))
    y += dh + s(14)

    # Ask Claude button
    bf  = font(fs(14), bold=True)
    bh  = s(46)
    rr(draw, [pad, y, pad + inner, y + bh], s(10), C["sky500"])
    btxt = "Ask Claude"
    draw.text(((cw - tw(draw, btxt, bf)) // 2, y + s(14)), btxt, font=bf, fill=rgb(C["white"]))
    y += bh + s(12)

    if show_answer:
        cyf  = font(fs(10), bold=True)
        qf   = font(fs(12))
        ansf = font(fs(13))
        lines = [
            "Yes, bring an umbrella! Toronto is currently",
            "experiencing light rain with 85% humidity",
            "and 14°C. Skies remain overcast through the",
            "afternoon — a waterproof jacket is wise.",
        ]
        bh2 = s(14) + s(20) + s(22) + len(lines) * s(21) + s(12)
        rr(draw, [pad, y, pad + inner, y + bh2], s(12), C["green50"], C["green300"])
        yi = y + s(12)
        draw.text((pad + s(12), yi), "TORONTO", font=cyf, fill=rgb(C["green600"]))
        yi += s(20)
        draw.text((pad + s(12), yi), '"Should I bring an umbrella?"', font=qf, fill=rgb(C["gray600"]))
        yi += s(22)
        for line in lines:
            draw.text((pad + s(12), yi), line, font=ansf, fill=rgb(C["gray900"]))
            yi += s(21)
        y += bh2 + s(14)

    # Interactive group box
    all_btns = ["MCP", "RAG", "Agent", "Claude", "History"]
    badf  = font(fs(10), bold=True)
    cap_f = font(fs(10), bold=True)
    btn_h = s(34)
    gpad  = s(8)
    gap   = s(6)
    n     = len(all_btns)
    col_w = (inner - gpad * 2 - gap * (n - 1)) // n
    gb_h  = s(20) + btn_h + s(14)
    rr(draw, [pad, y + s(10), pad + inner, y + s(10) + gb_h], s(10), C["bg"], C["gray300"])
    cap_txt = "Interactive"
    cap_w   = tw(draw, cap_txt, cap_f) + s(6)
    draw.rectangle([pad + s(10), y + s(4), pad + s(10) + cap_w, y + s(18)], fill=rgb(C["bg"]))
    draw.text((pad + s(13), y + s(5)), cap_txt, font=cap_f, fill=rgb(C["slate500"]))
    bx = pad + gpad
    by = y + s(10) + s(14)
    for badge in all_btns:
        border = C["sky300"]  if badge == "History" else C["violet200"]
        txt_c  = C["sky600"]  if badge == "History" else C["violet700"]
        rr(draw, [bx, by, bx + col_w, by + btn_h], s(8), C["white"], border)
        lx = bx + (col_w - tw(draw, badge, badf)) // 2
        draw.text((lx, by + s(10)), badge, font=badf, fill=rgb(txt_c))
        bx += col_w + gap
    y = by + btn_h + s(20)

    return img.crop((0, 0, cw, y))


def render_app(vw: int, show_answer: bool,
               max_content_w: int = None,
               viewport_h: int = None) -> Image.Image:
    """
    Render app at canvas width vw.
    max_content_w: simulate Tailwind max-w-* centering (content narrower than canvas).
    viewport_h: crop to this height (simulates above-the-fold for landscape).
    """
    cw      = min(vw, max_content_w) if max_content_w else vw
    content = _render_content(cw, show_answer)

    if cw < vw:
        # Paste content centered on full-width background
        xoff   = (vw - cw) // 2
        canvas = Image.new("RGB", (vw, content.height), rgb(C["bg"]))
        canvas.paste(content, (xoff, 0))
        content = canvas

    if viewport_h and viewport_h < content.height:
        content = content.crop((0, 0, vw, viewport_h))

    return content


# ══════════════════════════════════════════════════════════════════════════════
# DEVICE FRAME WRAPPERS
# ══════════════════════════════════════════════════════════════════════════════

def _rounded_mask(size, radius):
    mask = Image.new("L", size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, size[0], size[1]], radius=radius, fill=255)
    return mask


def ios_frame(screen: Image.Image) -> Image.Image:
    sw, sh = screen.size
    bw, bt, bb = 14, 56, 24   # bezel: sides, top, bottom
    hi        = 28             # home indicator zone
    fw        = sw + bw * 2
    fh        = sh + bt + bb + hi

    frame = Image.new("RGB", (fw, fh), "#1C1C1E")
    draw  = ImageDraw.Draw(frame)

    # Screen area (paste content, clip to screen shape)
    screen_mask = _rounded_mask((sw, sh), 6)
    frame.paste(screen, (bw, bt), screen_mask)

    # Dynamic island
    diw, dih = int(sw * 0.28), 26
    dix = (fw - diw) // 2
    diy = (bt - dih) // 2
    draw.rounded_rectangle([dix, diy, dix + diw, diy + dih], radius=13, fill="#000000")

    # Home indicator
    hiw = int(fw * 0.33)
    hix = (fw - hiw) // 2
    hiy = sh + bt + bb + (hi - 5) // 2
    draw.rounded_rectangle([hix, hiy, hix + hiw, hiy + 5], radius=3, fill="#555555")

    # Side buttons (volume + power)
    draw.rounded_rectangle([0, bt + 40, 3, bt + 80],  radius=2, fill="#3A3A3C")  # vol up
    draw.rounded_rectangle([0, bt + 90, 3, bt + 130], radius=2, fill="#3A3A3C")  # vol down
    draw.rounded_rectangle([fw - 3, bt + 60, fw, bt + 130], radius=2, fill="#3A3A3C")  # power

    # Rounded outer frame
    result = Image.new("RGB", (fw, fh), "#FFFFFF")
    result.paste(frame, mask=_rounded_mask((fw, fh), 42))
    return result


def android_frame(screen: Image.Image) -> Image.Image:
    sw, sh = screen.size
    bw, bt, bb = 13, 44, 14
    nb        = 38             # nav bar height
    fw        = sw + bw * 2
    fh        = sh + bt + bb + nb

    frame = Image.new("RGB", (fw, fh), "#202124")
    draw  = ImageDraw.Draw(frame)

    # Screen
    screen_mask = _rounded_mask((sw, sh), 4)
    frame.paste(screen, (bw, bt), screen_mask)

    # Punch-hole camera
    cr = 9
    cx, cy = fw // 2, bt // 2
    draw.ellipse([cx - cr, cy - cr, cx + cr, cy + cr], fill="#111111")
    draw.ellipse([cx - 4, cy - 4, cx + 4, cy + 4], fill="#333333")

    # Nav bar — gesture pill
    pilw = int(fw * 0.25)
    pilx = (fw - pilw) // 2
    pily = sh + bt + bb + (nb - 5) // 2
    draw.rounded_rectangle([pilx, pily, pilx + pilw, pily + 5], radius=3, fill="#555555")

    # Power button (right side)
    draw.rounded_rectangle([fw - 3, bt + 50, fw, bt + 110], radius=2, fill="#3A3A3C")

    result = Image.new("RGB", (fw, fh), "#FFFFFF")
    result.paste(frame, mask=_rounded_mask((fw, fh), 38))
    return result


def ipad_frame(screen: Image.Image, landscape: bool = False) -> Image.Image:
    sw, sh  = screen.size
    bside   = 18   # narrow uniform bezel
    bcam    = 36   # wider edge that holds the camera

    if landscape:
        # camera bar on the left
        bL, bR, bT, bB = bcam, bside, bside, bside
    else:
        # camera bar on top
        bL, bR, bT, bB = bside, bside, bcam, bside

    fw = sw + bL + bR
    fh = sh + bT + bB

    frame = Image.new("RGB", (fw, fh), "#2C2C2E")
    draw  = ImageDraw.Draw(frame)

    # Screen
    screen_mask = _rounded_mask((sw, sh), 6)
    frame.paste(screen, (bL, bT), screen_mask)

    # Camera dot on the centre of the camera-bar edge
    cr = 7
    if landscape:
        cx, cy = bL // 2, fh // 2
    else:
        cx, cy = fw // 2, bT // 2
    draw.ellipse([cx - cr, cy - cr, cx + cr, cy + cr], fill="#111111")
    draw.ellipse([cx - 3,  cy - 3,  cx + 3,  cy + 3],  fill="#333333")

    # Side / top-bottom hardware buttons
    if landscape:
        # power on top, volume on right
        draw.rounded_rectangle([bL + 40, 0, bL + 80,  3],  radius=2, fill="#3A3A3C")
        draw.rounded_rectangle([fw - 3,  bT + 30, fw, bT + 70],  radius=2, fill="#3A3A3C")
        draw.rounded_rectangle([fw - 3,  bT + 80, fw, bT + 120], radius=2, fill="#3A3A3C")
    else:
        # power on right, volume on left
        draw.rounded_rectangle([fw - 3, bT + 50, fw, bT + 110], radius=2, fill="#3A3A3C")
        draw.rounded_rectangle([0, bT + 40, 3, bT + 80],  radius=2, fill="#3A3A3C")
        draw.rounded_rectangle([0, bT + 90, 3, bT + 130], radius=2, fill="#3A3A3C")

    result = Image.new("RGB", (fw, fh), "#FFFFFF")
    result.paste(frame, mask=_rounded_mask((fw, fh), 34))
    return result


def web_frame(screen: Image.Image) -> Image.Image:
    sw, sh    = screen.size
    chrome_h  = 58
    tab_h     = 34
    total_h_c = chrome_h + tab_h
    fw        = sw + 2
    fh        = sh + total_h_c + 1

    frame = Image.new("RGB", (fw, fh), "#DEE1E6")
    draw  = ImageDraw.Draw(frame)

    # Tab bar
    draw.rectangle([0, 0, fw, tab_h], fill="#E8EAED")
    draw.rounded_rectangle([8, 6, 180, tab_h - 2], radius=6, fill="#FFFFFF")
    tf = font(11)
    draw.text((18, 14), "Weather AI", font=tf, fill="#3C4043")
    # tab close x
    draw.text((162, 12), "×", font=font(13), fill="#9AA0A6")

    # Toolbar
    draw.rectangle([0, tab_h, fw, tab_h + chrome_h], fill="#F1F3F4")

    # Traffic lights
    for i, color in enumerate(["#FF5F57", "#FEBC2E", "#28C840"]):
        cx = 18 + i * 22
        cy = tab_h + chrome_h // 2
        draw.ellipse([cx - 7, cy - 7, cx + 7, cy + 7], fill=color)

    # Back / forward
    nav_f = font(16)
    draw.text((68, tab_h + 16), "←", font=nav_f, fill="#AAAAAA")
    draw.text((88, tab_h + 16), "→", font=nav_f, fill="#CCCCCC")

    # Address bar
    ab_x0, ab_y0 = 110, tab_h + 10
    ab_x1, ab_y1 = fw - 60, tab_h + chrome_h - 10
    draw.rounded_rectangle([ab_x0, ab_y0, ab_x1, ab_y1], radius=14, fill="#FFFFFF",
                            outline="#DADCE0", width=1)
    url_f = font(12)
    draw.text((ab_x0 + 14, ab_y0 + 9), "localhost:8081  —  Weather AI", font=url_f, fill="#444746")

    # Reload icon
    draw.text((fw - 48, tab_h + 17), "↻", font=font(16), fill="#5F6368")

    # Separator
    draw.line([0, tab_h + chrome_h, fw, tab_h + chrome_h], fill="#DADCE0", width=1)

    # Screen content
    frame.paste(screen, (1, total_h_c))

    # Outer border
    draw.rectangle([0, 0, fw - 1, fh - 1], outline="#DADCE0", width=1)

    return frame


# ══════════════════════════════════════════════════════════════════════════════
# SETTINGS PAGE RENDERER
# ══════════════════════════════════════════════════════════════════════════════

def _render_settings_content(cw: int) -> Image.Image:
    scale = cw / 390
    def s(v): return max(1, int(v * scale))
    def fs(v): return max(8, int(v * scale))

    img   = Image.new("RGB", (cw, 3200), rgb(C["bg"]))
    draw  = ImageDraw.Draw(img)
    pad   = s(20)
    inner = cw - pad * 2
    y     = 0

    # ── Page header bar ──────────────────────────────────────────────────────
    hbar = s(80)
    draw.rectangle([0, 0, cw, hbar], fill=rgb(C["bg"]))
    draw.line([0, hbar - 1, cw, hbar - 1], fill=rgb(C["sky200"]), width=1)
    draw.text((s(18), s(44)), "☰", font=font(fs(20)),      fill=rgb(C["sky700"]))
    ttl_f = font(fs(18), bold=True)
    ttxt  = "Settings"
    draw.text(((cw - tw(draw, ttxt, ttl_f)) // 2, s(46)), ttxt, font=ttl_f, fill=rgb(C["sky700"]))
    y = hbar + s(14)

    sec_f = font(fs(10), bold=True)
    lbl_f = font(fs(14))
    hnt_f = font(fs(11))
    ctl_f = font(fs(12))

    def draw_switch(sx, sy, on):
        sw_w, sw_h = s(44), s(24)
        rr(draw, [sx, sy, sx + sw_w, sy + sw_h], sw_h // 2,
           C["sky500"] if on else C["gray300"])
        tx = sx + sw_w - sw_h + 3 if on else sx + 3
        draw.ellipse([tx, sy + 3, tx + sw_h - 6, sy + sw_h - 3], fill="#FFFFFF")

    def draw_section(title_txt, items, yp):
        row_h  = s(54)
        card_h = s(26) + row_h * len(items)
        rr(draw, [pad, yp, pad + inner, yp + card_h], s(12), C["white"])
        draw.text((pad + s(14), yp + s(7)), title_txt, font=sec_f, fill=rgb(C["slate500"]))
        ry = yp + s(26)
        for i, item in enumerate(items):
            if i < len(items) - 1:
                draw.line([pad + s(14), ry + row_h, pad + inner - s(14), ry + row_h],
                          fill=rgb(C["gray300"]), width=1)
            draw.text((pad + s(16), ry + s(9)), item["label"], font=lbl_f, fill=rgb(C["gray900"]))
            if item.get("hint"):
                draw.text((pad + s(16), ry + s(27)), item["hint"], font=hnt_f, fill=rgb(C["gray400"]))
            rx   = pad + inner - s(14)
            mid  = ry + row_h // 2
            ctrl = item.get("ctrl")
            if ctrl in ("on", "off"):
                draw_switch(rx - s(44), mid - s(12), ctrl == "on")
            elif ctrl == "picker":
                pw, ph = s(116), s(28)
                px, py = rx - pw, mid - ph // 2
                rr(draw, [px, py, px + pw, py + ph], s(6), C["white"], C["gray300"])
                draw.text((px + s(8), py + s(6)), item.get("val", "System"), font=ctl_f, fill=rgb(C["gray900"]))
                draw.text((px + pw - s(18), py + s(6)), "▾", font=ctl_f, fill=rgb(C["gray400"]))
            elif ctrl == "input":
                pw, ph = s(106), s(28)
                px, py = rx - pw, mid - ph // 2
                rr(draw, [px, py, px + pw, py + ph], s(8), C["white"], C["gray300"])
                val = item.get("val", "")
                draw.text((px + s(8), py + s(6)), val or "e.g. Toronto",
                          font=ctl_f, fill=rgb(C["gray900"] if val else C["gray400"]))
            ry += row_h
        return yp + card_h + s(12)

    y = draw_section("RUNNING MODE", [
        {"label": "Online  (ONLINE=1)", "hint": "Makes live API calls to fetch weather data", "ctrl": "on"},
    ], y)
    y = draw_section("STATE SYNC", [
        {"label": "Disabled  (STATE_SYNC=0)", "hint": "Persisted state is preserved across restarts", "ctrl": "off"},
    ], y)
    y = draw_section("LANGUAGE MODEL", [
        {"label": "LLM Provider", "hint": "AI model used by the Weather chatbot", "ctrl": "picker", "val": "Server Default"},
    ], y)
    y = draw_section("APPEARANCE", [
        {"label": "Theme", "ctrl": "picker", "val": "System"},
    ], y)
    y = draw_section("NAVIGATION", [
        {"label": "Default page", "hint": "Page shown on app launch and after login", "ctrl": "picker", "val": "Home"},
    ], y)
    y = draw_section("NOTIFICATIONS", [
        {"label": "Enable notifications", "hint": "Receive weather alerts and updates", "ctrl": "on"},
    ], y)
    y = draw_section("WEATHER", [
        {"label": "Default city", "hint": "Pre-fills the city field in Weather AI", "ctrl": "input", "val": "Toronto"},
    ], y)

    # ── Typography section (4 pickers + live preview) ─────────────────────────
    typo_items = [
        {"label": "Font Family", "ctrl": "picker", "val": "System"},
        {"label": "Font Size",   "ctrl": "picker", "val": "Medium"},
        {"label": "Font Weight", "ctrl": "picker", "val": "Regular"},
        {"label": "Font Style",  "ctrl": "picker", "val": "Normal"},
    ]
    row_h  = s(54)
    prev_h = s(52)
    card_h = s(26) + row_h * len(typo_items) + prev_h
    rr(draw, [pad, y, pad + inner, y + card_h], s(12), C["white"])
    draw.text((pad + s(14), y + s(7)), "TYPOGRAPHY", font=sec_f, fill=rgb(C["slate500"]))
    ry = y + s(26)
    for i, item in enumerate(typo_items):
        draw.line([pad + s(14), ry + row_h, pad + inner - s(14), ry + row_h],
                  fill=rgb(C["gray300"]), width=1)
        draw.text((pad + s(16), ry + s(9)), item["label"], font=lbl_f, fill=rgb(C["gray900"]))
        pw, ph = s(116), s(28)
        px_ = pad + inner - s(14) - pw
        py_ = ry + row_h // 2 - ph // 2
        rr(draw, [px_, py_, px_ + pw, py_ + ph], s(6), C["white"], C["gray300"])
        draw.text((px_ + s(8), py_ + s(6)), item["val"], font=ctl_f, fill=rgb(C["gray900"]))
        draw.text((px_ + pw - s(18), py_ + s(6)), "▾",  font=ctl_f, fill=rgb(C["gray400"]))
        ry += row_h
    # preview box
    draw.line([pad + s(14), ry, pad + inner - s(14), ry], fill=rgb(C["gray300"]), width=1)
    pf = font(fs(14))
    draw.text((pad + s(16), ry + s(10)),
              "The quick brown fox jumps over the lazy dog.",
              font=pf, fill=rgb(C["gray700"]))
    draw.text((pad + s(16), ry + s(28)), "Aa Bb Cc Dd  0 1 2 3",
              font=pf, fill=rgb(C["gray700"]))
    y += card_h + s(12)

    # ── Buttons ───────────────────────────────────────────────────────────────
    bf   = font(fs(14), bold=True)
    bh_b = s(44)
    for label in ("Default Settings", "Clear Settings"):
        rr(draw, [pad, y, pad + inner, y + bh_b], s(10), C["white"], C["sky300"], ow=2)
        draw.text(((cw - tw(draw, label, bf)) // 2, y + s(13)), label, font=bf, fill=rgb(C["sky600"]))
        y += bh_b + s(10)

    y += s(10)
    return img.crop((0, 0, cw, y))


def render_settings_page(vw: int, max_content_w: int = None) -> Image.Image:
    cw      = min(vw, max_content_w) if max_content_w else vw
    content = _render_settings_content(cw)
    if cw < vw:
        canvas = Image.new("RGB", (vw, content.height), rgb(C["bg"]))
        canvas.paste(content, ((vw - cw) // 2, 0))
        content = canvas
    return content


# ══════════════════════════════════════════════════════════════════════════════
# SYNC-STATES MODAL RENDERER
# ══════════════════════════════════════════════════════════════════════════════

def _render_sync_modal_content(cw: int) -> Image.Image:
    scale = cw / 390
    def s(v): return max(1, int(v * scale))
    def fs(v): return max(8, int(v * scale))

    # Settings page as dimmed background
    bg_full = _render_settings_content(cw)
    crop_h  = min(bg_full.height, s(680))
    bg      = bg_full.crop((0, 0, cw, crop_h))
    overlay = Image.new("RGBA", (cw, crop_h), (0, 0, 0, 120))
    result  = Image.alpha_composite(bg.convert("RGBA"), overlay).convert("RGB")
    draw    = ImageDraw.Draw(result)

    pad   = s(20)
    inner = cw - pad * 2

    # Modal sheet rising from bottom
    modal_h = s(496)
    my      = crop_h - modal_h
    rr(draw, [0, my, cw, crop_h + s(30)], s(20), C["white"])

    # Title
    draw.text((pad, my + s(18)), "Sync State", font=font(fs(16), bold=True), fill=rgb(C["gray900"]))

    # Description
    df = font(fs(12))
    draw.text((pad, my + s(44)), "No .env state vars found. Paste a JSON object", font=df, fill=rgb(C["gray600"]))
    draw.text((pad, my + s(60)), "below to seed the app state.", font=df, fill=rgb(C["gray600"]))

    # Default Settings button (outline, sky)
    dbf   = font(fs(13), bold=True)
    db_y  = my + s(80)
    db_h  = s(38)
    rr(draw, [pad, db_y, pad + inner, db_y + db_h], s(10), C["white"], C["sky300"], ow=2)
    draw.text(((cw - tw(draw, "Default Settings", dbf)) // 2, db_y + s(10)),
              "Default Settings", font=dbf, fill=rgb(C["sky600"]))

    # Clear Settings button (outline, red)
    cb_y = db_y + db_h + s(8)
    cb_h = s(38)
    rr(draw, [pad, cb_y, pad + inner, cb_y + cb_h], s(10), C["white"], "#FCA5A5", ow=2)
    draw.text(((cw - tw(draw, "Clear Settings", dbf)) // 2, cb_y + s(10)),
              "Clear Settings", font=dbf, fill=rgb("#DC2626"))

    # JSON textarea
    ta_y = cb_y + cb_h + s(10)
    ta_h = s(170)
    rr(draw, [pad, ta_y, pad + inner, ta_y + ta_h], s(10), "#F9FAFB", C["gray300"])
    jf = font(fs(10))
    json_lines = [
        '{',
        '  "user": { "email": "demo@example.com", "phone": "" },',
        '  "settings": {',
        '    "theme": "system", "defaultCity": "Toronto",',
        '    "online": true, "stateSync": true',
        '  },',
        '  "weather": {',
        '    "city": "Toronto",',
        '    "question": "Should I bring an umbrella?"',
        '  },',
        '  "questions": { "items": ["Should I bring...", ...] }',
        '}',
    ]
    jy = ta_y + s(10)
    for line in json_lines:
        draw.text((pad + s(10), jy), line, font=jf, fill=rgb(C["gray700"]))
        jy += s(17)

    # Buttons
    btn_y = ta_y + ta_h + s(14)
    btn_w = (inner - s(10)) // 2
    btn_h = s(42)
    cbf   = font(fs(13), bold=True)

    for i, label in enumerate(("Cancel", "Apply")):
        bx = pad + i * (btn_w + s(10))
        rr(draw, [bx, btn_y, bx + btn_w, btn_y + btn_h], s(10), C["sky500"])
        draw.text(((bx + (btn_w - tw(draw, label, cbf)) // 2), btn_y + s(12)),
                  label, font=cbf, fill=rgb(C["white"]))

    return result


# ══════════════════════════════════════════════════════════════════════════════
# HISTORY PAGE RENDERER
# ══════════════════════════════════════════════════════════════════════════════

def _render_history_content(cw: int) -> Image.Image:
    scale = cw / 390
    def s(v): return max(1, int(v * scale))
    def fs(v): return max(8, int(v * scale))

    img  = Image.new("RGB", (cw, 2400), rgb(C["bg"]))
    draw = ImageDraw.Draw(img)
    y    = 0

    # ── Page header bar ──────────────────────────────────────────────────────
    hbar = s(80)
    draw.rectangle([0, 0, cw, hbar], fill=rgb(C["bg"]))
    draw.line([0, hbar - 1, cw, hbar - 1], fill=rgb(C["sky200"]), width=1)
    draw.text((s(18), s(44)), "☰", font=font(fs(20)), fill=rgb(C["sky700"]))
    ttl_f  = font(fs(18), bold=True)
    ttxt   = "History"
    draw.text(((cw - tw(draw, ttxt, ttl_f)) // 2, s(46)), ttxt, font=ttl_f, fill=rgb(C["sky700"]))
    cl_f   = font(fs(13), bold=True)
    cl_txt = "Clear"
    draw.text((cw - s(14) - tw(draw, cl_txt, cl_f), s(48)), cl_txt, font=cl_f, fill=rgb("#DC2626"))
    y = hbar

    # ── Export toolbar ────────────────────────────────────────────────────────
    tb_h = s(46)
    draw.rectangle([0, y, cw, y + tb_h], fill=rgb(C["white"]))
    draw.line([0, y + tb_h - 1, cw, y + tb_h - 1], fill=rgb(C["gray300"]), width=1)
    lbl_f2 = font(fs(12))
    draw.text((s(14), y + s(15)), "Export as", font=lbl_f2, fill=rgb(C["gray400"]))
    lbl_w = tw(draw, "Export as", lbl_f2) + s(8)
    # picker box
    pk_x, pk_y = s(14) + lbl_w, y + s(8)
    pk_w, pk_h = s(148), s(30)
    rr(draw, [pk_x, pk_y, pk_x + pk_w, pk_y + pk_h], s(6), C["white"], C["gray300"])
    pk_f = font(fs(11))
    draw.text((pk_x + s(8), pk_y + s(8)), "SQL  (UNION SELECT)", font=pk_f, fill=rgb(C["gray900"]))
    draw.text((pk_x + pk_w - s(18), pk_y + s(8)), "▾", font=pk_f, fill=rgb(C["gray400"]))
    # export button
    ex_x = pk_x + pk_w + s(10)
    ex_w = s(72)
    ex_h = pk_h
    rr(draw, [ex_x, pk_y, ex_x + ex_w, pk_y + ex_h], s(7), C["sky500"])
    ex_f = font(fs(12), bold=True)
    ex_txt = "Export"
    draw.text((ex_x + (ex_w - tw(draw, ex_txt, ex_f)) // 2, pk_y + s(8)),
              ex_txt, font=ex_f, fill=rgb(C["white"]))
    y += tb_h

    sh_f    = font(fs(11), bold=True)
    time_f  = font(fs(11))
    label_f = font(fs(12))
    mono_f  = font(fs(9))
    row_h   = s(42)

    def draw_date_header(date_str, yp):
        sh_h = s(32)
        draw.rectangle([0, yp, cw, yp + sh_h], fill=rgb(C["sky50"]))
        draw.line([0, yp + sh_h - 1, cw, yp + sh_h - 1], fill=rgb(C["sky200"]), width=1)
        draw.text((s(16), yp + s(9)), date_str, font=sh_f, fill=rgb(C["sky700"]))
        return yp + sh_h

    def draw_row(t_str, label, yp, expanded_json=None):
        extra_h = 0
        json_lines = []
        if expanded_json:
            json_lines = expanded_json
            extra_h = s(10) + len(json_lines) * s(15) + s(6)
        rh = row_h + extra_h
        draw.rectangle([0, yp, cw, yp + rh], fill=rgb(C["white"]))
        draw.line([0, yp + rh - 1, cw, yp + rh - 1], fill=rgb("#F3F4F6"), width=1)
        draw.text((s(16),        yp + s(14)), t_str, font=time_f,  fill=rgb(C["gray400"]))
        draw.text((s(16) + s(72), yp + s(13)), label, font=label_f, fill=rgb(C["gray900"]))
        if json_lines:
            jy = yp + row_h + s(4)
            for line in json_lines:
                draw.text((s(16) + s(72), jy), line, font=mono_f, fill=rgb(C["gray600"]))
                jy += s(15)
        return yp + rh

    # ── Today ────────────────────────────────────────────────────────────────
    y = draw_date_header("2026-05-24", y)
    y = draw_row("09:15:32", "State synced from .env", y)
    y = draw_row("09:15:33", "Navigate → weather-ai", y)
    y = draw_row("09:16:01", "Weather: Toronto — Should I bring an umbrella?", y,
                 expanded_json=[
                     '{"type": "weather_ask",',
                     ' "label": "Weather: Toronto...",',
                     ' "data": {',
                     '   "city": "Toronto",',
                     '   "question": "Should I bring an umbrella?"',
                     ' }',
                     '}',
                 ])
    y = draw_row("09:16:09", 'Setting: LlmProvider → "claude"', y)
    y = draw_row("09:17:44", "Weather: Paris — Will it rain this weekend?", y)
    y = draw_row("09:18:03", "User logged out", y)

    # ── Yesterday ────────────────────────────────────────────────────────────
    y += s(4)
    y = draw_date_header("2026-05-23", y)
    y = draw_row("16:42:11", "Navigate → settings", y)
    y = draw_row("16:42:18", 'Setting: Theme → "dark"', y)
    y = draw_row("16:43:05", "State synced from .env", y)
    y = draw_row("16:55:30", "User: demo@example.com", y)

    y += s(16)
    return img.crop((0, 0, cw, y))


def render_history_page(vw: int, max_content_w: int = None) -> Image.Image:
    cw      = min(vw, max_content_w) if max_content_w else vw
    content = _render_history_content(cw)
    if cw < vw:
        canvas = Image.new("RGB", (vw, content.height), rgb(C["bg"]))
        canvas.paste(content, ((vw - cw) // 2, 0))
        content = canvas
    return content


# ══════════════════════════════════════════════════════════════════════════════
# SAVE SCREENSHOTS
# ══════════════════════════════════════════════════════════════════════════════
def save_screenshots():
    mob_app       = render_app(390, show_answer=True)
    web_app       = render_app(860, show_answer=True)
    ipad_port_app = render_app(540,  show_answer=True, max_content_w=512)
    ipad_land_app = render_app(1000, show_answer=True, max_content_w=512)

    settings_mob  = render_settings_page(390)
    settings_web  = render_settings_page(860, max_content_w=512)
    modal_content = _render_sync_modal_content(390)
    history_mob   = render_history_page(390)
    history_web   = render_history_page(860, max_content_w=512)

    ios_img          = ios_frame(mob_app)
    android_img      = android_frame(mob_app)
    web_img          = web_frame(web_app)
    ipad_port_img    = ipad_frame(ipad_port_app, landscape=False)
    ipad_land_img    = ipad_frame(ipad_land_app, landscape=True)
    settings_ios_img = ios_frame(settings_mob)
    settings_web_img = web_frame(settings_web)
    modal_ios_img    = ios_frame(modal_content)
    history_ios_img  = ios_frame(history_mob)
    history_web_img  = web_frame(history_web)

    ios_path             = ROOT / "preview_ios.jpg"
    android_path         = ROOT / "preview_android.jpg"
    web_path             = ROOT / "preview_web.jpg"
    ipad_port_path       = ROOT / "preview_ipad_portrait.jpg"
    ipad_land_path       = ROOT / "preview_ipad_landscape.jpg"
    settings_ios_path    = ROOT / "preview_settings_ios.jpg"
    settings_web_path    = ROOT / "preview_settings_web.jpg"
    modal_path           = ROOT / "preview_sync_modal.jpg"
    history_ios_path     = ROOT / "preview_history_ios.jpg"
    history_web_path     = ROOT / "preview_history_web.jpg"

    ios_img.save(str(ios_path),                "JPEG", quality=93)
    android_img.save(str(android_path),        "JPEG", quality=93)
    web_img.save(str(web_path),                "JPEG", quality=93)
    ipad_port_img.save(str(ipad_port_path),    "JPEG", quality=93)
    ipad_land_img.save(str(ipad_land_path),    "JPEG", quality=93)
    settings_ios_img.save(str(settings_ios_path), "JPEG", quality=93)
    settings_web_img.save(str(settings_web_path), "JPEG", quality=93)
    modal_ios_img.save(str(modal_path),        "JPEG", quality=93)
    history_ios_img.save(str(history_ios_path), "JPEG", quality=93)
    history_web_img.save(str(history_web_path), "JPEG", quality=93)

    for img, p in [
        (ios_img, ios_path), (android_img, android_path), (web_img, web_path),
        (ipad_port_img, ipad_port_path), (ipad_land_img, ipad_land_path),
        (settings_ios_img, settings_ios_path), (settings_web_img, settings_web_path),
        (modal_ios_img, modal_path),
        (history_ios_img, history_ios_path), (history_web_img, history_web_path),
    ]:
        print(f"saved: {p.name:<36} ({img.size[0]}×{img.size[1]})")

    return (ios_path, android_path, web_path, ipad_port_path, ipad_land_path,
            settings_ios_path, settings_web_path, modal_path,
            history_ios_path, history_web_path)


# ══════════════════════════════════════════════════════════════════════════════
# EMBED INTO DOCX
# ══════════════════════════════════════════════════════════════════════════════
def _caption(cell, text):
    p = cell.add_paragraph(text)
    p.runs[0].font.size = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def embed_into_docx(ios_path, android_path, web_path, ipad_port_path, ipad_land_path,
                    settings_ios_path, settings_web_path, modal_path,
                    history_ios_path, history_web_path):
    doc  = Document(str(DOCX))
    body = doc.element.body
    NS   = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    # Remove everything from the first "UI Preview" heading onward
    cut_from = None
    for i, child in enumerate(body):
        if child.tag == f"{{{NS}}}p":
            text = "".join(t.text or "" for t in child.iter(f"{{{NS}}}t"))
            if "UI Preview" in text:
                cut_from = i - 1
                break

    if cut_from is not None:
        sect_pr = body.find(f"{{{NS}}}sectPr")
        for child in list(body)[cut_from:]:
            if child is not sect_pr:
                body.remove(child)

    # ── Page 1: iOS · Android · Web ──────────────────────────────────────────
    doc.add_page_break()
    h = doc.add_heading("UI Preview — iOS · Android · Web", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.add_run(f"Generated: {datetime.date.today()}").font.size = Pt(9)
    doc.add_paragraph()

    t1 = doc.add_table(rows=1, cols=3)
    t1.style = "Table Grid"

    t1.cell(0, 0).paragraphs[0].add_run().add_picture(str(ios_path),     width=Inches(1.9))
    _caption(t1.cell(0, 0), "iOS — iPhone 14")

    t1.cell(0, 1).paragraphs[0].add_run().add_picture(str(android_path), width=Inches(1.9))
    _caption(t1.cell(0, 1), "Android — Pixel")

    t1.cell(0, 2).paragraphs[0].add_run().add_picture(str(web_path),     width=Inches(2.4))
    _caption(t1.cell(0, 2), "Web — Browser")

    # ── Page 2: iPad Portrait · iPad Landscape ───────────────────────────────
    doc.add_page_break()
    h2 = doc.add_heading("UI Preview — iPad", level=1)
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()

    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Table Grid"

    t2.cell(0, 0).paragraphs[0].add_run().add_picture(str(ipad_port_path), width=Inches(1.8))
    _caption(t2.cell(0, 0), "iPad — Portrait (full page)")

    t2.cell(0, 1).paragraphs[0].add_run().add_picture(str(ipad_land_path), width=Inches(3.6))
    _caption(t2.cell(0, 1), "iPad — Landscape (full page)")

    # ── Page 3: Settings · Modal ─────────────────────────────────────────────
    doc.add_page_break()
    h3 = doc.add_heading("UI Preview — Settings & Sync State Modal", level=1)
    h3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()

    t3 = doc.add_table(rows=1, cols=3)
    t3.style = "Table Grid"

    t3.cell(0, 0).paragraphs[0].add_run().add_picture(str(settings_ios_path), width=Inches(1.9))
    _caption(t3.cell(0, 0), "Settings — iOS")

    t3.cell(0, 1).paragraphs[0].add_run().add_picture(str(settings_web_path), width=Inches(2.4))
    _caption(t3.cell(0, 1), "Settings — Web")

    t3.cell(0, 2).paragraphs[0].add_run().add_picture(str(modal_path), width=Inches(1.9))
    _caption(t3.cell(0, 2), "Sync State Modal — iOS")

    # ── Page 4: History Page ─────────────────────────────────────────────────
    doc.add_page_break()
    h4 = doc.add_heading("UI Preview — History Page", level=1)
    h4.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()

    t4 = doc.add_table(rows=1, cols=2)
    t4.style = "Table Grid"

    t4.cell(0, 0).paragraphs[0].add_run().add_picture(str(history_ios_path), width=Inches(1.9))
    _caption(t4.cell(0, 0), "History — iOS")

    t4.cell(0, 1).paragraphs[0].add_run().add_picture(str(history_web_path), width=Inches(2.4))
    _caption(t4.cell(0, 1), "History — Web")

    doc.save(str(DOCX))
    print(f"docx updated: {DOCX.name}")


if __name__ == "__main__":
    ios_p, and_p, web_p, ipad_port_p, ipad_land_p, set_ios_p, set_web_p, modal_p, hist_ios_p, hist_web_p = save_screenshots()
    embed_into_docx(ios_p, and_p, web_p, ipad_port_p, ipad_land_p, set_ios_p, set_web_p, modal_p, hist_ios_p, hist_web_p)
